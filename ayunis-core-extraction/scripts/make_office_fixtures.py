"""Generates the Office fixtures the conversion tests use, and the large acceptance document.

    uv run python -m scripts.make_office_fixtures            # small committed set
    uv run python -m scripts.make_office_fixtures large DIR  # 200-page DOCX, ignored by git

Each small fixture carries the same municipal prose as the PDF set, so the text Gotenberg and
pypdfium2 produce can be compared against `<name>.expected.txt` next to it.
"""

import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from pptx import Presentation
from pptx.util import Inches

from scripts.make_fixtures import PROSE

FIXTURE_DIR = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "office"

ODF_MANIFEST = """<?xml version="1.0" encoding="UTF-8"?>
<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"
 manifest:version="1.2">
 <manifest:file-entry manifest:full-path="/" manifest:media-type="{mime}"/>
 <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
</manifest:manifest>
"""

ODF_CONTENT = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
 xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
 xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
 xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0"
 xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0"
 xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0"
 office:version="1.2">
 <office:automatic-styles>
  <style:style style:name="dp1" style:family="drawing-page">
   <style:drawing-page-properties draw:fill="none"/>
  </style:style>
  <style:style style:name="gr1" style:family="graphic">
   <style:graphic-properties draw:fill="none" draw:stroke="none"/>
  </style:style>
 </office:automatic-styles>
 <office:body>{body}</office:body>
</office:document-content>
"""


def paragraphs(count: int) -> list[str]:
    return [PROSE[index % len(PROSE)] for index in range(count)]


def docx_fixture(path: Path, pages: int = 2, paragraphs_per_page: int = 6) -> None:
    document = Document()
    for page in range(pages):
        document.add_heading(f"Beschlussvorlage {page + 1}", level=1)
        for line in paragraphs(paragraphs_per_page):
            document.add_paragraph(line)
        if page < pages - 1:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.save(str(path))


def pptx_fixture(path: Path, slides: int = 2) -> None:
    presentation = Presentation()
    layout = presentation.slide_layouts[1]
    for index in range(slides):
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = f"Sitzung des Gemeinderats, Folie {index + 1}"
        body = slide.placeholders[1].text_frame
        body.text = PROSE[0]
        for line in PROSE[1:3]:
            body.add_paragraph().text = line
        slide.shapes.add_textbox(
            Inches(1), Inches(5.5), Inches(8), Inches(1)
        ).text_frame.text = PROSE[3]
    presentation.save(str(path))


def _odf(path: Path, mime: str, body: str) -> None:
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(zipfile.ZipInfo("mimetype"), mime, compress_type=zipfile.ZIP_STORED)
        archive.writestr("META-INF/manifest.xml", ODF_MANIFEST.format(mime=mime))
        archive.writestr("content.xml", ODF_CONTENT.format(body=body))


def odt_fixture(path: Path) -> None:
    lines = "".join(f"<text:p>{line}</text:p>" for line in paragraphs(8))
    _odf(
        path,
        "application/vnd.oasis.opendocument.text",
        f'<office:text><text:h text:outline-level="1">Satzung</text:h>{lines}</office:text>',
    )


def odp_fixture(path: Path) -> None:
    pages = "".join(
        f'<draw:page draw:name="page{index + 1}" draw:style-name="dp1">'
        '<draw:frame draw:style-name="gr1" svg:width="24cm" svg:height="12cm"'
        ' svg:x="2cm" svg:y="3cm">'
        f"<draw:text-box><text:p>Folie {index + 1}</text:p>"
        + "".join(f"<text:p>{line}</text:p>" for line in PROSE[:3])
        + "</draw:text-box></draw:frame></draw:page>"
        for index in range(2)
    )
    _odf(
        path,
        "application/vnd.oasis.opendocument.presentation",
        f"<office:presentation>{pages}</office:presentation>",
    )


def rtf_fixture(path: Path) -> None:
    def escape(text: str) -> str:
        return "".join(ch if ord(ch) < 128 else f"\\u{ord(ch)}?" for ch in text)

    body = "\\par\n".join(escape(line) for line in paragraphs(8))
    path.write_text(
        "{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Helvetica;}}\\f0\\fs22\n"
        "{\\b Bescheid}\\par\n" + body + "\\par\n}\n",
        encoding="ascii",
    )


def small_set(directory: Path = FIXTURE_DIR) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    docx_fixture(directory / "letter.docx")
    pptx_fixture(directory / "slides.pptx")
    odt_fixture(directory / "letter.odt")
    odp_fixture(directory / "slides.odp")
    rtf_fixture(directory / "letter.rtf")


def large_set(directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    docx_fixture(directory / "letter-200.docx", pages=200, paragraphs_per_page=8)


if __name__ == "__main__":
    if len(sys.argv) >= 3 and sys.argv[1] == "large":
        large_set(Path(sys.argv[2]))
    else:
        small_set()
